from __future__ import annotations

import json
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from mutagen import File as MutagenFile
from sklearn.feature_extraction.text import TfidfVectorizer
from wordcloud import WordCloud

from .utils import AUDIO_EXTS, TRANSCRIPT_EXTS, guess_pair_key


@dataclass
class FileRecord:
    file_id: str
    audio_rel: Optional[str]
    transcript_rel: Optional[str]
    duration_sec: Optional[float]
    word_count: int
    top_terms: List[Tuple[str, float]]
    # placeholders for next modules
    label: Optional[str] = None
    label_scores: Optional[Dict[str, float]] = None
    summary: Optional[str] = None


def _read_docx_text(path: Path) -> str:
    """
    Read DOCX text robustly.
    - First try python-docx (best for paragraphs)
    - If DOCX is partially corrupted (common: media CRC errors), fallback to parsing word/document.xml only.
    """
    try:
        from docx import Document  # local import to avoid hard fail
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        return "\n".join(parts)
    except Exception:
        # fallback parse document.xml
        try:
            import zipfile
            import xml.etree.ElementTree as ET
            from io import BytesIO

            data = path.read_bytes()
            with zipfile.ZipFile(BytesIO(data)) as zf:
                xml_bytes = zf.read("word/document.xml")

            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            root = ET.fromstring(xml_bytes)
            texts = [t.text for t in root.findall(".//w:t", ns) if t.text]
            return " ".join(texts)
        except Exception:
            return ""


def _read_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".docx":
        return _read_docx_text(path)
    if ext == ".json":
        try:
            obj = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
            if isinstance(obj, dict):
                for k in ("text", "transcript", "full_text"):
                    if k in obj and isinstance(obj[k], str):
                        return obj[k]
                return json.dumps(obj, ensure_ascii=False, indent=2)
            if isinstance(obj, list):
                texts: List[str] = []
                for item in obj:
                    if isinstance(item, dict) and isinstance(item.get("text"), str):
                        texts.append(item["text"])
                if texts:
                    return "\n".join(texts)
                return json.dumps(obj, ensure_ascii=False, indent=2)
        except Exception:
            return path.read_text(encoding="utf-8", errors="ignore")
    return ""


def _audio_duration_seconds(path: Path) -> Optional[float]:
    try:
        mf = MutagenFile(str(path))
        if mf is None or mf.info is None:
            return None
        length = getattr(mf.info, "length", None)
        if length is None:
            return None
        return float(length)
    except Exception:
        return None


def _tokenize_text(text: str) -> str:
    # very light cleanup for TF-IDF/wordcloud
    import re

    t = text.lower()
    t = re.sub(r"http\S+|www\.\S+", " ", t)
    t = re.sub(r"[^0-9a-zA-ZÀ-ÿ\u0100-\u017F\u0180-\u024F\u1E00-\u1EFF\s]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def _tfidf_top_terms(corpus: List[str], top_k: int = 12) -> List[Tuple[str, float]]:
    if not corpus:
        return []
    vec = TfidfVectorizer(
        max_features=2000,
        ngram_range=(1, 2),
        stop_words=None,
    )
    X = vec.fit_transform(corpus)

    import numpy as np

    scores = np.asarray(X.mean(axis=0)).ravel()
    terms = vec.get_feature_names_out()
    idx = scores.argsort()[::-1][:top_k]
    return [(terms[i], float(scores[i])) for i in idx if scores[i] > 0]


def _make_wordcloud(text: str, out_path: Path, *, max_words: int = 200, max_chars: int = 200000) -> None:
    if not text.strip():
        return
    text = (text or "")[:max_chars]
    wc = WordCloud(
        width=1200,
        height=600,
        background_color="white",
        collocations=False,
        max_words=max_words,
    ).generate(text)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wc.to_file(str(out_path))


def analyze_job(job_dir: Path) -> Dict:
    """
    job_dir/
      extracted/   (input data after unzip)
      analysis.json
      artifacts/   (wordclouds)
    """
    extracted_dir = job_dir / "extracted"
    artifacts_dir = job_dir / "artifacts"
    artifacts_dir.mkdir(parents=True, exist_ok=True)

    audio_files: Dict[str, Path] = {}
    transcript_files: Dict[str, Path] = {}

    for p in extracted_dir.rglob("*"):
        if not p.is_file():
            continue
        ext = p.suffix.lower()
        key = guess_pair_key(p)
        if ext in AUDIO_EXTS:
            audio_files[key] = p
        elif ext in TRANSCRIPT_EXTS:
            transcript_files[key] = p

    keys = sorted(set(audio_files.keys()) | set(transcript_files.keys()))
    records: List[FileRecord] = []
    texts_for_global: List[str] = []

    for key in keys:
        audio_p = audio_files.get(key)
        tr_p = transcript_files.get(key)

        audio_rel = str(audio_p.relative_to(extracted_dir)).replace("\\", "/") if audio_p else None
        tr_rel = str(tr_p.relative_to(extracted_dir)).replace("\\", "/") if tr_p else None

        duration = _audio_duration_seconds(audio_p) if audio_p else None
        raw_text = _read_text(tr_p) if tr_p else ""
        clean = _tokenize_text(raw_text)

        word_count = len(clean.split()) if clean else 0
        top_terms = _tfidf_top_terms([clean], top_k=10) if clean else []

        if clean:
            texts_for_global.append(clean)

        records.append(
            FileRecord(
                file_id=key,
                audio_rel=audio_rel,
                transcript_rel=tr_rel,
                duration_sec=duration,
                word_count=word_count,
                top_terms=top_terms,
                label=None,
                label_scores=None,
                summary=None,
            )
        )

    global_text = " ".join(texts_for_global)
    global_top = _tfidf_top_terms(texts_for_global, top_k=16) if texts_for_global else []
    if global_text:
        _make_wordcloud(global_text, artifacts_dir / "wc_global.png", max_words=250, max_chars=250000)

    total_duration = sum([r.duration_sec or 0 for r in records])
    total_words = sum([r.word_count for r in records])

    out: Dict = {
        "job_id": job_dir.name,
        "created_at": __import__("datetime").datetime.now().isoformat(timespec="seconds"),
        "counts": {
            "files_total": len(records),
            "files_with_audio": sum(1 for r in records if r.audio_rel),
            "files_with_transcript": sum(1 for r in records if r.transcript_rel),
        },
        "totals": {
            "duration_sec": float(total_duration),
            "word_count": int(total_words),
        },
        "global_top_terms": global_top,
        "records": [asdict(r) for r in records],
        "artifacts": {
            "wordcloud_global": "artifacts/wc_global.png" if (artifacts_dir / "wc_global.png").exists() else None
        },
    }

    (job_dir / "analysis.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    return out
